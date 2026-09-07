"""Unit tests for document parsers and ParserRegistry."""

from __future__ import annotations

import io

import docx
import pytest
from pypdf import PdfWriter

from rag_llm_services_rag.parsers.docx import DocxParser
from rag_llm_services_rag.parsers.markdown import MarkdownParser
from rag_llm_services_rag.parsers.pdf import PDFParser
from rag_llm_services_rag.parsers.registry import get_default_parser_registry
from rag_llm_services_rag.parsers.text import TextParser
from rag_llm_services_shared.errors import UnsupportedMediaTypeError, ValidationError

# -----------------------------------------------------------------------------
# TextParser tests
# -----------------------------------------------------------------------------


def test_text_parser_basic_utf8() -> None:
    parser = TextParser()
    assert "text/plain" in parser.supported_mime_types()
    assert ".txt" in parser.supported_extensions()

    content = "Xin chào Việt Nam!\nĐây là tài liệu thử nghiệm.".encode()
    doc = parser.parse(content, filename="test.txt")

    assert "Xin chào Việt Nam!" in doc.raw_text
    assert len(doc.sections) == 1
    assert doc.sections[0].page_number == 1
    assert doc.metadata["filename"] == "test.txt"


def test_text_parser_latin1_fallback() -> None:
    parser = TextParser()
    # 0xE9 is 'é' in latin-1, but invalid standalone byte in UTF-8
    content = b"Caf\xe9 au lait"
    doc = parser.parse(content)
    assert "Café au lait" in doc.raw_text or "Caf" in doc.raw_text


def test_text_parser_empty_content() -> None:
    parser = TextParser()
    doc = parser.parse(b"")
    assert doc.raw_text == ""
    assert len(doc.sections) == 1


# -----------------------------------------------------------------------------
# MarkdownParser tests
# -----------------------------------------------------------------------------


def test_markdown_parser_heading_hierarchy() -> None:
    parser = MarkdownParser()
    assert "text/markdown" in parser.supported_mime_types()
    assert ".md" in parser.supported_extensions()

    md_content = """# Tổng quan hệ thống
Tài liệu kiến trúc RAG LLM Services.

## Thành phần Backend
Mô tả các service FastAPI.

### Cơ sở dữ liệu
PostgreSQL kết hợp pgvector.

## Triển khai
Docker compose và Kubernetes.
"""
    doc = parser.parse(md_content.encode("utf-8"), filename="arch.md")

    assert len(doc.sections) == 4
    # Section 1: H1
    assert doc.sections[0].section_header == "Tổng quan hệ thống"
    assert doc.sections[0].hierarchy == ["Tổng quan hệ thống"]
    assert "Tài liệu kiến trúc" in doc.sections[0].content

    # Section 2: H2
    assert doc.sections[1].section_header == "Thành phần Backend"
    assert doc.sections[1].hierarchy == ["Tổng quan hệ thống", "Thành phần Backend"]

    # Section 3: H3
    assert doc.sections[2].section_header == "Cơ sở dữ liệu"
    assert doc.sections[2].hierarchy == [
        "Tổng quan hệ thống",
        "Thành phần Backend",
        "Cơ sở dữ liệu",
    ]

    # Section 4: H2 (popped H3)
    assert doc.sections[3].section_header == "Triển khai"
    assert doc.sections[3].hierarchy == ["Tổng quan hệ thống", "Triển khai"]


def test_markdown_parser_no_headings() -> None:
    parser = MarkdownParser()
    raw = "Văn bản thông thường không có tiêu đề markdown."
    doc = parser.parse(raw.encode("utf-8"))
    assert len(doc.sections) == 1
    assert doc.sections[0].section_header is None
    assert doc.sections[0].hierarchy == []
    assert doc.sections[0].content == raw


# -----------------------------------------------------------------------------
# PDFParser tests
# -----------------------------------------------------------------------------


def _create_test_pdf_bytes(page_count: int = 2) -> bytes:
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_pdf_parser_valid() -> None:
    parser = PDFParser()
    assert "application/pdf" in parser.supported_mime_types()
    assert ".pdf" in parser.supported_extensions()

    pdf_bytes = _create_test_pdf_bytes(page_count=3)
    doc = parser.parse(pdf_bytes, filename="sample.pdf")

    assert doc.metadata["mime_type"] == "application/pdf"
    assert doc.metadata["total_pages"] == 3


def test_pdf_parser_corrupted_raises_validation_error() -> None:
    parser = PDFParser()
    corrupt_bytes = b"%PDF-1.4\ncorrupted content that cannot be parsed %%EOF"
    with pytest.raises(ValidationError) as excinfo:
        parser.parse(corrupt_bytes)
    assert "Invalid or corrupted PDF" in str(excinfo.value)


# -----------------------------------------------------------------------------
# DocxParser tests
# -----------------------------------------------------------------------------


def _create_test_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("Chương 1: Mở đầu", level=1)
    doc.add_paragraph("Nội dung chương 1 mô tả về quy trình xử lý dữ liệu.")
    doc.add_heading("Phần 1.1: Thu thập", level=2)
    doc.add_paragraph("Thu thập văn bản từ nhiều nguồn khác nhau.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[1].cells[0].text = "Data 1"
    table.rows[1].cells[1].text = "Data 2"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_parser_valid() -> None:
    parser = DocxParser()
    docx_bytes = _create_test_docx_bytes()
    doc = parser.parse(docx_bytes, filename="document.docx")

    assert len(doc.sections) >= 2
    assert doc.sections[0].section_header == "Chương 1: Mở đầu"
    assert "quy trình xử lý dữ liệu" in doc.sections[0].content
    assert doc.sections[1].section_header == "Phần 1.1: Thu thập"
    assert doc.sections[1].hierarchy == ["Chương 1: Mở đầu", "Phần 1.1: Thu thập"]
    assert "Header A | Header B" in doc.raw_text


def test_docx_parser_corrupted_raises_validation_error() -> None:
    parser = DocxParser()
    corrupt_bytes = b"PK\x03\x04corrupted zip payload not a valid docx"
    with pytest.raises(ValidationError) as excinfo:
        parser.parse(corrupt_bytes)
    assert "Invalid or corrupted DOCX" in str(excinfo.value)


# -----------------------------------------------------------------------------
# ParserRegistry tests
# -----------------------------------------------------------------------------


def test_registry_lookup_by_mime_and_extension() -> None:
    registry = get_default_parser_registry()

    assert isinstance(registry.select_by_mime_and_sniff("application/pdf"), PDFParser)
    assert isinstance(registry.select_by_mime_and_sniff("text/plain"), TextParser)
    assert isinstance(registry.select_by_mime_and_sniff("text/markdown"), MarkdownParser)
    assert isinstance(registry.select_by_mime_and_sniff("text/x-markdown"), MarkdownParser)
    assert isinstance(
        registry.select_by_mime_and_sniff(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        DocxParser,
    )

    # Extension fallback
    assert isinstance(
        registry.select_by_mime_and_sniff("application/octet-stream", filename="manual.pdf"),
        PDFParser,
    )
    assert isinstance(
        registry.select_by_mime_and_sniff("application/octet-stream", filename="notes.docx"),
        DocxParser,
    )


def test_registry_content_sniffing() -> None:
    registry = get_default_parser_registry()

    pdf_header = b"%PDF-1.4\nSome PDF content"
    selected = registry.select_by_mime_and_sniff(
        "application/octet-stream", filename=None, content=pdf_header
    )
    assert isinstance(selected, PDFParser)


def test_registry_unsupported_media_type_raises() -> None:
    registry = get_default_parser_registry()
    with pytest.raises(UnsupportedMediaTypeError) as excinfo:
        registry.select_by_mime_and_sniff("video/mp4", filename="clip.mp4")
    assert "No document parser available" in str(excinfo.value)
